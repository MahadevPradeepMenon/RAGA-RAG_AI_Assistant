from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".pptx"}


def load_txt_file(file_path: Path) -> str:
    """
    Read a single .txt file and return its text.
    """
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()


def load_pdf_file(file_path: Path) -> str:
    """
    Read a PDF file and extract text from each page.
    Works best for text-based PDFs.
    """
    reader = PdfReader(str(file_path))
    text_parts = []

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text()

        if page_text:
            text_parts.append(f"[Page {page_number}]\n{page_text}")

    return "\n\n".join(text_parts)


def load_docx_file(file_path: Path) -> str:
    """
    Read a DOCX file and extract text from paragraphs and tables.
    """
    document = Document(str(file_path))
    text_parts = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text_parts.append(paragraph_text)

    for table_index, table in enumerate(document.tables, start=1):
        text_parts.append(f"\n[Table {table_index}]")

        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip().replace("\n", " ")

                if cell_text:
                    cells.append(cell_text)

            if cells:
                text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def extract_text_from_shape(shape) -> list[str]:
    """
    Extract text from a PowerPoint shape.

    Handles:
    - normal text boxes
    - placeholders
    - tables
    """
    text_parts = []

    if hasattr(shape, "text"):
        shape_text = shape.text.strip()

        if shape_text:
            text_parts.append(shape_text)

    if shape.has_table:
        for row in shape.table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip().replace("\n", " ")

                if cell_text:
                    cells.append(cell_text)

            if cells:
                text_parts.append(" | ".join(cells))

    return text_parts


def load_pptx_slides(file_path: Path) -> list[dict]:
    """
    Read a PPTX file and return one document per slide.

    This is slide-aware:
    each slide becomes its own searchable document with slide metadata.
    """
    presentation = Presentation(str(file_path))
    slide_documents = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text_parts = []

        for shape in slide.shapes:
            slide_text_parts.extend(extract_text_from_shape(shape))

        slide_text = "\n".join(slide_text_parts).strip()

        if slide_text:
            slide_documents.append({
                "source": file_path.name,
                "text": slide_text,
                "file_type": ".pptx",
                "location": f"Slide {slide_number}",
            })

    return slide_documents


def load_single_document(file_path: Path) -> dict:
    """
    Load one supported non-PPTX document and return source + text.
    """
    extension = file_path.suffix.lower()

    if extension == ".txt":
        text = load_txt_file(file_path)
    elif extension == ".pdf":
        text = load_pdf_file(file_path)
    elif extension == ".docx":
        text = load_docx_file(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    return {
        "source": file_path.name,
        "text": text,
        "file_type": extension,
        "location": None,
    }


def load_documents(folder_path: str) -> list[dict]:
    """
    Load all supported documents from a folder.

    Supports:
    - .txt
    - .pdf
    - .docx
    - .pptx

    PPTX files are loaded slide-by-slide.
    """
    folder = Path(folder_path)
    documents = []

    if not folder.exists():
        raise FileNotFoundError(f"Folder not found: {folder_path}")

    for file_path in folder.iterdir():
        if not file_path.is_file():
            continue

        extension = file_path.suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            continue

        if extension == ".pptx":
            slide_documents = load_pptx_slides(file_path)
            documents.extend(slide_documents)
        else:
            document = load_single_document(file_path)

            if document["text"].strip():
                documents.append(document)

    return documents